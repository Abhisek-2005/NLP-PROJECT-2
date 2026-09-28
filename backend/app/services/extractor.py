import io
import logging
from pathlib import Path
from fastapi import HTTPException
import fitz  # PyMuPDF
import docx

logger = logging.getLogger(__name__)

class DocumentExtractorService:
    """Extracts plain text from various resume file formats (PDF, DOCX, TXT)."""

    @staticmethod
    def extract_text_from_pdf(content: bytes) -> str:
        """Extracts text from PDF bytes using PyMuPDF (fitz)."""
        try:
            doc = fitz.open(stream=content, filetype="pdf")
            if doc.is_encrypted:
                raise ValueError("PDF is password-protected.")

            text_chunks = []
            for page_num in range(len(doc)):
                page = doc[page_num]
                text = page.get_text("text")
                if text:
                    text_chunks.append(text.strip())
            
            full_text = "\n\n".join(text_chunks).strip()
            if not full_text:
                raise ValueError("PDF does not contain extractable text (it may be a scanned image).")
            return full_text
        except Exception as e:
            logger.error("Failed to parse PDF: %s", str(e))
            raise HTTPException(
                status_code=400,
                detail=f"Failed to extract text from PDF: {str(e)}"
            )

    @staticmethod
    def extract_text_from_docx(content: bytes) -> str:
        """Extracts text from DOCX bytes using python-docx."""
        try:
            file_stream = io.BytesIO(content)
            doc = docx.Document(file_stream)
            
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            
            # Also extract text from any tables in the docx
            table_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        table_text.append(" | ".join(row_data))
                        
            full_text = "\n".join(paragraphs + table_text).strip()
            if not full_text:
                raise ValueError("DOCX document contains no readable text.")
            return full_text
        except Exception as e:
            logger.error("Failed to parse DOCX: %s", str(e))
            raise HTTPException(
                status_code=400,
                detail=f"Failed to extract text from DOCX: {str(e)}"
            )

    @staticmethod
    def extract_text_from_txt(content: bytes) -> str:
        """Decodes plain text bytes with fallback encodings."""
        for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                return content.decode(enc).strip()
            except UnicodeDecodeError:
                continue
        raise HTTPException(status_code=400, detail="Unable to decode text file. Ensure it is UTF-8 encoded.")

    @classmethod
    def extract(cls, filename: str, content: bytes) -> str:
        """Route to appropriate extractor based on file extension."""
        ext = Path(filename).suffix.lower()
        if not content:
            raise HTTPException(status_code=400, detail="Uploaded file is empty.")

        if ext == ".pdf":
            return cls.extract_text_from_pdf(content)
        elif ext in (".docx", ".doc"):
            if ext == ".doc":
                raise HTTPException(
                    status_code=400,
                    detail="Legacy .doc format is not supported. Please convert to modern .docx or .pdf."
                )
            return cls.extract_text_from_docx(content)
        elif ext in (".txt", ".md", ".rtf"):
            return cls.extract_text_from_txt(content)
        else:
            raise HTTPException(
                status_code=400,
                detail=f"Unsupported file format '{ext}'. Please upload a PDF (.pdf), Word document (.docx), or Text file (.txt)."
            )
